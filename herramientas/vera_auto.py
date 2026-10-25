#!/usr/bin/env python3
"""VERA · generador de informes automáticos (sin revisar).

Uso: python3 vera_auto.py datos.json salida_sin_extension

Entrada: JSON con los datos extraídos del EPAR o del ensayo, cada cifra con su página.
Salida: salida.docx y salida.pdf (convertido con LibreOffice).

Principios de VERA que aplica:
- La extracción automática solo extrae cifras; todos los cálculos (RAR, NNT, NNH, LHH, interacción de subgrupos,
  puntuación de Gil-Sierra) se hacen aquí con fórmulas.
- Cada dato lleva su página y queda «sin verificar».
- Aviso de responsabilidad en la portada y en el pie de cada página.
"""
import json, math, os, subprocess, sys, tempfile

import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

AZUL = RGBColor(0x1F, 0x4E, 0x79)
GRIS = RGBColor(0x5B, 0x66, 0x75)
ROJO = RGBColor(0xA3, 0x30, 0x2A)
TINTA = RGBColor(0x1B, 0x24, 0x30)
Z = 1.959964

AVISO_CORTO = ('Informe generado automáticamente a partir de los documentos aportados y de fuentes públicas. '
               'No revisado por VERA: verificar los datos es responsabilidad del usuario.')


# ---------------------------------------------------------------- cálculos
def fmt(x, d=1):
    s = f'{x:,.{d}f}'.replace(',', 'X').replace('.', ',').replace('X', '.')
    return s.replace('-', '−')


def wald_diff(e1, n1, e2, n2):
    """Diferencia de riesgos p2 − p1 (control − tratamiento) con IC 95 % de Wald."""
    p1, p2 = e1 / n1, e2 / n2
    d = p2 - p1
    se = math.sqrt(p1 * (1 - p1) / n1 + p2 * (1 - p2) / n2)
    return d, d - Z * se, d + Z * se


def nnt_text(d, lo, hi):
    """NNT (o NNH) a partir de la diferencia de riesgos y su IC (en proporción)."""
    if d == 0:
        return '—', 'No calculable'
    n = math.ceil(1 / abs(d))
    if lo <= 0 <= hi:
        return str(n), 'IC incluye el 0: no significativo'
    a, b = sorted([math.ceil(1 / abs(lo)), math.ceil(1 / abs(hi))])
    return str(n), f'IC 95 % {a} a {b}'


def altman_nnt(pc_evt, hr):
    """NNT de Altman (riesgos proporcionales): 1 / (Sc^HR − Sc), con Sc = supervivencia en el control."""
    sc = 1 - pc_evt
    den = sc ** hr - sc
    return None if den <= 0 else math.ceil(1 / den)


def interaction_q(rows):
    """Q de Cochran sobre log(HR) de los subgrupos; devuelve (Q, gl, p)."""
    ys, ws = [], []
    for hr, lo, hi in rows:
        se = (math.log(hi) - math.log(lo)) / (2 * Z)
        ys.append(math.log(hr)); ws.append(1 / se ** 2)
    ybar = sum(w * y for w, y in zip(ws, ys)) / sum(ws)
    q = sum(w * (y - ybar) ** 2 for w, y in zip(ws, ys))
    df = len(rows) - 1
    return q, df, chi2_sf(q, df)


def chi2_sf(x, k):
    """Probabilidad de cola superior de la chi cuadrado (series exactas para k pequeño)."""
    if k == 1:
        return math.erfc(math.sqrt(x / 2))
    if k == 2:
        return math.exp(-x / 2)
    # k > 2: integración por series de la gamma incompleta regularizada
    a, xx = k / 2, x / 2
    s, term, n = 0.0, 1 / a, 0
    while term > 1e-12 and n < 500:
        s += term; n += 1; term *= xx / (a + n)
    p_lower = s * math.exp(-xx + a * math.log(xx) - math.lgamma(a))
    return max(0.0, 1 - p_lower)


def bucher_md(ab, cb):
    """Comparación indirecta de Bucher para diferencias de medias con comparador común.
    ab = A frente a C, cb = B frente a C, cada uno [estimación, li, ls] (IC 95 %). Devuelve A frente a B."""
    se = lambda e: (e[2] - e[1]) / (2 * Z)
    d = ab[0] - cb[0]
    s = math.sqrt(se(ab) ** 2 + se(cb) ** 2)
    p = math.erfc(abs(d / s) / math.sqrt(2))
    return d, d - Z * s, d + Z * s, p


def bucher_ratio(ab, cb):
    """Bucher para HR/RR/OR en escala logarítmica."""
    se = lambda e: (math.log(e[2]) - math.log(e[1])) / (2 * Z)
    d = math.log(ab[0]) - math.log(cb[0])
    s = math.sqrt(se(ab) ** 2 + se(cb) ** 2)
    p = math.erfc(abs(d / s) / math.sqrt(2))
    return math.exp(d), math.exp(d - Z * s), math.exp(d + Z * s), p


GS_PUNTOS = {'Probable': 3, 'Posible': 2, 'Dudosa': 0, 'Nula': -3}


def gil_sierra(puntos):
    """Aplica la regla del cuestionario: asociación o consistencia nula → aplicabilidad nula."""
    val = {k: v for k, v, _ in puntos}
    total = sum(GS_PUNTOS[v] for _, v, _ in puntos)
    if val.get('Asociación') == 'Nula' or val.get('Consistencia') == 'Nula':
        return total, 'NULA'
    if total >= 7: return total, 'PROBABLE'
    if total >= 5: return total, 'POSIBLE'
    if total >= 3: return total, 'DUDOSA'
    return total, 'NULA'


# ---------------------------------------------------------------- Word
def set_cell_bg(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), hex_color)
    tcPr.append(shd)


def set_borders(table, color='D9DDE3'):
    tbl = table._tbl
    tblPr = tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        el = OxmlElement(f'w:{edge}')
        el.set(qn('w:val'), 'single'); el.set(qn('w:sz'), '4'); el.set(qn('w:space'), '0'); el.set(qn('w:color'), color)
        borders.append(el)
    tblPr.append(borders)


def add_field(run, instr):
    for kind, text in (('begin', None), (None, instr), ('separate', None), (None, '1'), ('end', None)):
        if kind:
            fc = OxmlElement('w:fldChar'); fc.set(qn('w:fldCharType'), kind); run._r.append(fc)
        elif text == instr:
            it = OxmlElement('w:instrText'); it.set(qn('xml:space'), 'preserve'); it.text = instr; run._r.append(it)
        else:
            t = OxmlElement('w:t'); t.text = text; run._r.append(t)


class Informe:
    def __init__(self, data):
        self.d = data
        self.doc = Document()
        st = self.doc.styles['Normal']
        st.font.name = 'Calibri'; st.font.size = Pt(10); st.font.color.rgb = TINTA
        st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Calibri')
        st.paragraph_format.space_after = Pt(4)
        st.paragraph_format.line_spacing = 1.15
        for name, size in (('Heading 1', 15), ('Heading 2', 12), ('Heading 3', 10.5)):
            h = self.doc.styles[name]
            h.font.name = 'Georgia' if name != 'Heading 3' else 'Calibri'
            h.font.size = Pt(size); h.font.bold = True; h.font.color.rgb = AZUL if name != 'Heading 3' else TINTA
            h.paragraph_format.space_before = Pt(14 if name == 'Heading 1' else 10)
            h.paragraph_format.space_after = Pt(4)
            h.paragraph_format.keep_with_next = True
        sec = self.doc.sections[0]
        sec.page_width, sec.page_height = Cm(21), Cm(29.7)
        sec.left_margin = sec.right_margin = Cm(2)
        sec.top_margin, sec.bottom_margin = Cm(2), Cm(2)
        self._header_footer(sec)

    # -- cabecera y pie en todas las páginas
    def _header_footer(self, sec):
        hp = sec.header.paragraphs[0]
        r = hp.add_run('VERA · Informe automático · '); r.font.size = Pt(8); r.font.color.rgb = GRIS
        r = hp.add_run('NO REVISADO'); r.font.size = Pt(8); r.bold = True; r.font.color.rgb = ROJO
        r = hp.add_run(f"   ·   {self.d['titulo']}"); r.font.size = Pt(8); r.font.color.rgb = GRIS
        sec.different_first_page_header_footer = True
        ff = sec.first_page_footer.paragraphs[0]
        r = ff.add_run(AVISO_CORTO); r.font.size = Pt(7.5); r.font.color.rgb = GRIS
        fp = sec.footer.paragraphs[0]
        r = fp.add_run(AVISO_CORTO + '  ·  Página '); r.font.size = Pt(7.5); r.font.color.rgb = GRIS
        r = fp.add_run(); r.font.size = Pt(7.5); r.font.color.rgb = GRIS; add_field(r, 'PAGE')
        r = fp.add_run(' de '); r.font.size = Pt(7.5); r.font.color.rgb = GRIS
        r = fp.add_run(); r.font.size = Pt(7.5); r.font.color.rgb = GRIS; add_field(r, 'NUMPAGES')

    # -- utilidades
    def p(self, text='', bold=False, color=None, size=None, italic=False, align=None):
        par = self.doc.add_paragraph()
        if text:
            r = par.add_run(text); r.bold = bold; r.italic = italic
            if color: r.font.color.rgb = color
            if size: r.font.size = Pt(size)
        if align: par.alignment = align
        return par

    def nota(self, text):
        return self.p(text, color=GRIS, size=8.5)

    def h(self, text, lvl=1):
        return self.doc.add_heading(text, lvl)

    def tabla(self, head, rows, widths=None, head_bg='E3ECF5', size=9):
        t = self.doc.add_table(rows=1, cols=len(head))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_borders(t)
        for i, txt in enumerate(head):
            c = t.rows[0].cells[i]; c.text = ''
            r = c.paragraphs[0].add_run(txt); r.bold = True; r.font.size = Pt(size - 0.5); r.font.color.rgb = AZUL
            set_cell_bg(c, head_bg)
        for row in rows:
            cells = t.add_row().cells
            for i, val in enumerate(row):
                cells[i].text = ''
                par = cells[i].paragraphs[0]
                if isinstance(val, tuple):  # (texto, nota pequeña)
                    r = par.add_run(val[0]); r.font.size = Pt(size)
                    r2 = par.add_run('\n' + val[1]); r2.font.size = Pt(size - 1.5); r2.font.color.rgb = GRIS
                else:
                    r = par.add_run(str(val)); r.font.size = Pt(size)
        tr = t.rows[0]._tr; trPr = tr.get_or_add_trPr()
        th = OxmlElement('w:tblHeader'); th.set(qn('w:val'), 'true'); trPr.append(th)
        if widths:
            for row in t.rows:
                for i, w in enumerate(widths):
                    row.cells[i].width = Cm(w)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)
        return t

    def caja(self, titulo, texto, bg='F8E3E1', color=ROJO):
        t = self.doc.add_table(rows=1, cols=1); set_borders(t, 'A3302A' if color == ROJO else 'D9DDE3')
        c = t.rows[0].cells[0]; set_cell_bg(c, bg); c.text = ''
        r = c.paragraphs[0].add_run(titulo); r.bold = True; r.font.color.rgb = color; r.font.size = Pt(10.5)
        for line in texto:
            pr = c.add_paragraph(); rr = pr.add_run(line); rr.font.size = Pt(9)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)

    def pg(self, pags):
        return f"{self.d['solicitud'].get('fuente_corta', 'EPAR')} p. {pags} · sin verificar"

    # -- secciones
    def portada(self):
        s = self.d['solicitud']
        logo = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'logo.png')
        sp = self.p(); sp.paragraph_format.space_before = Pt(36)
        if os.path.exists(logo):
            self.doc.add_picture(logo, width=Cm(3.2))
            self.doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        par = self.doc.add_paragraph(); par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = par.add_run('VERA'); r.font.name = 'Georgia'; r.font.size = Pt(54); r.bold = True; r.font.color.rgb = AZUL
        par.paragraph_format.space_after = Pt(0)
        self.p('Valoración de la Evidencia y Relevancia Asistencial', color=GRIS, size=12, align=WD_ALIGN_PARAGRAPH.CENTER)
        sp = self.p(); sp.paragraph_format.space_before = Pt(22)
        tp = self.doc.add_paragraph(); tp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = tp.add_run(self.d['titulo']); r.font.name = 'Georgia'; r.font.size = Pt(22); r.bold = True; r.font.color.rgb = TINTA
        self.p('Informe de evaluación generado automáticamente', color=GRIS, size=11, align=WD_ALIGN_PARAGRAPH.CENTER)
        sp = self.p(); sp.paragraph_format.space_before = Pt(12)
        for m in self.d['meta']:
            self.p(m, color=GRIS, size=9.5, align=WD_ALIGN_PARAGRAPH.CENTER)
        self.p(f"Solicitud del {s['fecha']} · Documentos aportados: {s['documentos']}", color=GRIS, size=9.5, align=WD_ALIGN_PARAGRAPH.CENTER)
        sp = self.p(); sp.paragraph_format.space_before = Pt(14)
        self.caja('Informe automático: no revisado por VERA', [
            'Este informe se ha generado automáticamente a partir de los documentos que aportaste y de fuentes externas públicas, señaladas como tales. Nadie del equipo de VERA ha comprobado sus datos.',
            'Cada cifra indica su fuente y su página y está marcada como «sin verificar». Los cálculos (RAR, NNT, NNH, interacción de subgrupos, comparación indirecta) se hacen con fórmulas a partir de esas cifras: si una cifra es errónea, el cálculo también lo será.',
            'Los apartados marcados como «propuesta automática» son valoraciones generadas automáticamente, no del evaluador.',
            'Verificar los datos antes de usar este informe es responsabilidad tuya. No sustituye el criterio de la Comisión de Farmacia y Terapéutica.'])
        self.doc.add_page_break()
        self.h('Ficha en 30 segundos', 2)
        self.tabla(['', 'Resultado', 'Detalle'], [[a, b, c] for a, b, c in self.d['ficha']], [4, 5, 8])

    def identificacion(self):
        i = self.d['identificacion']
        self.h('1. Identificación del fármaco e indicación')
        self.p(i['farmaco']); self.p('Posología: ' + i['posologia']); self.p('Indicación: ' + i['indicacion'])
        self.nota(self.pg(i['pag']))
        e = self.d.get('espana')
        if e:
            self.h('Situación en España', 3)
            self.tabla(['Fuente', 'Situación', 'Detalle', 'Consulta'], e['filas'], [3.2, 2.3, 7.5, 4])
            self.nota(e['nota'])
        else:
            self.nota('Situación en España (BIFIMED, IPT de la AEMPS): no consultada. Compruébala antes de usar el informe.')

    def ensayo(self):
        self.h('2. Ensayo pivotal: diseño y población')
        self.tabla(['Característica', 'Descripción', 'Fuente'], [[a, b, self.pg(c)] for a, b, c in self.d['diseno']], [3.5, 10, 3.5])
        self.h('Características basales', 3)
        self.tabla(['Variable', 'Valor', 'Fuente'], [[a, b, self.pg(c)] for a, b, c in self.d['basales']], [5.5, 8, 3.5])

    def eficacia(self):
        self.h('3. Eficacia')
        armT, armC = 'Nerandomilast 18 mg', 'Placebo'
        rows = [[(e['var'], self.pg(e['pag'])), e['t'], e['c'], e['dif']] for e in self.d['eficacia']]
        self.tabla(['Variable', armT, armC, 'Diferencia / HR (IC 95 %)'], rows, [6, 3, 3, 5])
        self.nota(self.d['eficacia_nota'])
        self.h('Beneficio absoluto (calculado)', 2)
        rows = []
        self._nnt_ref = None
        for b in self.d['beneficio']:
            d, lo, hi = wald_diff(b['et'], b['nt'], b['ec'], b['nc'])
            n, nci = nnt_text(d, lo, hi)
            alt = altman_nnt(b['ec'] / b['nc'], b['hr'][0])
            hr_sig = b['hr'][2] < 1 or b['hr'][1] > 1
            if not (lo <= 0 <= hi) and not hr_sig:
                nci += '. Diferencia bruta significativa pero HR no significativo: interpretar con cautela'
            rows.append([(b['lab'], self.pg(b['pag'])),
                         f"{fmt(100*b['et']/b['nt'])} % frente a {fmt(100*b['ec']/b['nc'])} %",
                         (f"{fmt(100*d)} % ({fmt(100*lo)} a {fmt(100*hi)})", 'IC 95 % de Wald'),
                         (n, nci),
                         (str(alt) if alt else '—', 'HR significativo' if hr_sig else 'HR no significativo')])
            if self._nnt_ref is None:
                self._nnt_ref = (b['lab'], d, lo, hi)
        self.tabla(['Variable (tiempo)', '% con el evento (18 mg / placebo)', 'Reducción absoluta', 'NNT', 'NNT (HR, Altman)'], rows, [5, 3.2, 3.6, 2.6, 2.6])
        self.nota('Cálculo propio a partir de los recuentos del documento: proporciones brutas durante todo el seguimiento (no estimaciones de Kaplan-Meier), sin ajustar por los factores de estratificación. El NNT de Altman usa el HR y supone riesgos proporcionales. Si el IC de la diferencia incluye el 0, el NNT no es significativo.')

    def subgrupos(self):
        s = self.d['subgrupos']
        self.h('4. Subgrupos')
        self.p(s['titulo'], bold=True)
        rows = s['filas']
        # forest plot
        fig, ax = plt.subplots(figsize=(6.6, 0.45 * (len(rows) + 1) + 0.6), dpi=200)
        labels = [r[0].strip() for r in rows] + ['Global']
        data = [r[2] for r in rows] + [s['global']]
        for k, (hr, lo, hi) in enumerate(data):
            y = len(data) - k
            col = '#1f4e79' if k < len(rows) else '#1b2430'
            ax.plot([lo, hi], [y, y], color=col, lw=1.4)
            ax.plot(hr, y, 's' if k < len(rows) else 'D', color=col, ms=6)
            ax.text(5.6, y, f"{fmt(hr,2)} ({fmt(lo,2)}–{fmt(hi,2)})", va='center', fontsize=8)
        ax.axvline(1, color='#888', lw=0.8, ls='--')
        ax.set_xscale('log'); ax.set_xlim(0.05, 5); ax.set_xticks([0.1, 0.25, 0.5, 1, 2, 4]); ax.set_xticklabels(['0,1', '0,25', '0,5', '1', '2', '4'], fontsize=8)
        ax.set_yticks(range(len(data), 0, -1)); ax.set_yticklabels(labels, fontsize=8.5)
        ax.set_xlabel('HR (escala logarítmica) · <1 favorece a nerandomilast', fontsize=8)
        for sp in ('top', 'right'): ax.spines[sp].set_visible(False)
        fig.tight_layout()
        img = os.path.join(self.tmp, 'forest.png'); fig.savefig(img); plt.close(fig)
        self.doc.add_picture(img, width=Cm(16))
        a, b = s['comparar']
        q, df, pq = interaction_q([rows[a][2], rows[b][2]])
        tab = [[(r[0].strip(), ''), r[1], f"{fmt(r[2][0],2)} ({fmt(r[2][1],2)} a {fmt(r[2][2],2)})"] for r in rows]
        self.tabla(['Subgrupo', 'Eventos 18 mg · placebo', 'HR (IC 95 %)'], tab, [5, 6, 6])
        self.p(f"Interacción calculada ({rows[a][0].strip()} frente a {rows[b][0].strip().lower()}): Q = {fmt(q,2)}, p = {fmt(pq,3)}. "
               f"Interacción en el documento: p = {s['p_epar']}.")
        self.nota(self.pg(s['pag']) + '. La p calculada usa la Q de Cochran sobre los HR publicados y puede diferir de la del modelo del ensayo.')
        if s.get('otros'):
            self.p(s['otros'])
        for g in self.d['gil_sierra']:
            self.h('Cuestionario de aplicabilidad de subgrupos (Gil-Sierra et al., 2020): ' + g['factor'], 3)
            self.p('Propuesta automática, pendiente del juicio del evaluador.', italic=True, color=ROJO, size=9)
            for x in g['preliminares']:
                self.p('• ' + x, size=9)
            self.tabla(['Criterio', 'Cuestión', 'Respuesta', 'Fiabilidad', 'Justificación'], g['items'], [2.8, 3, 3.6, 2.2, 5.4], size=8.5)
            total, ver = gil_sierra(g['puntos'])
            pts = [[k, v, str(GS_PUNTOS[v])] for k, v, _ in g['puntos']] + [['Aplicabilidad (calculada)', ver, str(total)]]
            self.tabla(['Criterio', 'Valoración', 'Puntos'], pts, [6, 6, 3])
            if ver != g['veredicto']:
                self.p(f"Aviso: el veredicto propuesto ({g['veredicto']}) no coincide con el calculado ({ver}). Se usa el calculado.", color=ROJO, size=9)
            self.p('Recomendación (' + ver.lower() + '): ' + g['recomendacion'])

    def seguridad(self):
        s = self.d['seguridad']
        self.h('5. Seguridad')
        rows = []
        lab, dn, _, _ = self._nnt_ref
        for name, et, nt, ec, nc, pag in s['filas']:
            d, lo, hi = wald_diff(ec, nc, et, nt)   # tratamiento − control: >0 = más frecuente con el fármaco
            if d > 0:
                n, nci = nnt_text(d, lo, hi); tipo = 'NNH'
            else:
                n, nci = nnt_text(d, lo, hi); tipo = 'NNT para evitarlo'
            lhh = '—'
            if tipo == 'NNH' and not (lo <= 0 <= hi) and dn > 0:
                lhh = fmt((1 / d) / (1 / dn), 1)
            rows.append([(name, self.pg(pag)), f"{fmt(100*et/nt)} %", f"{fmt(100*ec/nc)} %",
                         (f"{'+' if d > 0 else ''}{fmt(100*d)} ({fmt(100*lo)} a {fmt(100*hi)})", ''), (f'{tipo} {n}', nci), lhh])
        self.tabla(['Efecto adverso', '18 mg', 'Placebo', 'Diferencia de riesgo, % (IC 95 %)', 'NNH / NNT', 'LHH'], rows, [4.4, 1.6, 1.6, 3.8, 3.6, 1.2], size=8.5)
        self.nota(s['nota'] + f' Diferencias con IC de Wald calculados a partir de los recuentos. LHH = NNH / NNT frente a «{lab}» (>1 indica que es más probable el beneficio que el daño).')
        self.p(s['texto'])

    def indirecta(self):
        ind = self.d['indirecta']
        self.h('7. Comparación indirecta')
        if isinstance(ind, str):
            self.p(ind); return
        u = ind['unidad']
        self.p(f"Método de Bucher con placebo como comparador común. Variable: {ind['variable']}. Diferencia positiva = menor pérdida con nerandomilast.")
        rows = []
        for key in ('ab', 'ab_alt'):
            if key not in ind: continue
            a = ind[key]
            d, lo, hi, p = bucher_md(a['est'], ind['cb']['est']) if ind['tipo'] == 'MD' else bucher_ratio(a['est'], ind['cb']['est'])
            sig = 'No significativa' if lo <= (0 if ind['tipo'] == 'MD' else 1) <= hi else 'Significativa'
            rows.append([(a['lab'], a['fuente']), f"{fmt(a['est'][0])} ({fmt(a['est'][1])} a {fmt(a['est'][2])}) {u}",
                         (f"{fmt(d)} {u} ({fmt(lo)} a {fmt(hi)})", f"p = {fmt(p,2)} · {sig}")])
        c = ind['cb']
        self.tabla(['Nerandomilast frente a placebo (A–C)', 'Efecto A–C', f"Nerandomilast frente a {ind['comparador']} (calculado)"], rows, [6.5, 4.5, 6])
        self.p(f"Comparador frente a placebo (B–C): {c['lab']}: {fmt(c['est'][0],2)} {u} ({fmt(c['est'][1],2)} a {fmt(c['est'][2],2)}). Fuente: {c['fuente']}.", size=9)
        self.p('Propuesta automática sobre la validez, pendiente del juicio del evaluador:', italic=True, color=ROJO, size=9)
        for n in ind['notas']:
            self.p('• ' + n, size=9)

    def resto(self):
        self.h('6. Relevancia clínica'); self.p(self.d['relevancia'])
        self.nota('La escala ESMO-MCBS solo se aplica a tratamientos oncológicos.')
        self.indirecta()
        e = self.d['economica']
        self.h('8. Área económica'); self.p(e['nota'])
        self.tabla(['Tratamiento', 'Pauta', 'Coste por ciclo (€)', 'Coste paciente-año (€)', 'Coste del tratamiento (€)'],
                   [[a, b, '', '', ''] for a, b in e['filas']] + [['Incremental', '', '', '', '']], [4.5, 2, 3.2, 3.6, 3.7])
        self.nota('Celdas vacías para que introduzcas los precios de tu centro. VERA no presenta precios de ejemplo como reales.')
        self.h('9. Posicionamiento'); self.p('Propuesta automática, pendiente del juicio del evaluador.', italic=True, color=ROJO, size=9)
        self.p(self.d['posicionamiento'])
        self.h('Fuentes', 2)
        for f in self.d['fuentes']:
            self.p('• ' + f, size=9)
        self.caja('Recordatorio', [AVISO_CORTO, 'Generado el ' + self.d['solicitud']['fecha'] + ' por VERA (vera-fharma.github.io/Vera).'])

    def build(self, out):
        with tempfile.TemporaryDirectory() as tmp:
            self.tmp = tmp
            self.portada(); self.identificacion(); self.ensayo(); self.eficacia(); self.subgrupos(); self.seguridad(); self.resto()
            self.doc.core_properties.title = 'VERA · ' + self.d['titulo'] + ' (informe automático, no revisado)'
            self.doc.core_properties.author = 'VERA (generado automáticamente)'
            self.doc.save(out + '.docx')
        outdir = os.path.dirname(os.path.abspath(out)) or '.'
        subprocess.run(['soffice', '--headless', '--convert-to', 'pdf', '--outdir', outdir, out + '.docx'],
                       check=True, capture_output=True, timeout=180)
        return out + '.docx', out + '.pdf'


if __name__ == '__main__':
    data = json.load(open(sys.argv[1], encoding='utf-8'))
    print(Informe(data).build(sys.argv[2]))
